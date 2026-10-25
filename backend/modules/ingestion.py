from pathlib import Path

import fitz  # PyMuPDF
from docx import Document

import config
from modules import vector_store
from utils.chunker import chunk_text
from utils.embeddings import embed_text


def ingest_document(
    file_path: str, doc_id: str, filename: str, chat_id: str
) -> dict:
    """Detect file type, parse, chunk, embed, store. Returns status dict."""
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        blocks = parse_pdf(file_path)
    elif suffix == ".docx":
        blocks = parse_docx(file_path)
    elif suffix == ".txt":
        blocks = parse_txt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    all_chunks, all_embeddings, all_metadatas, all_ids = [], [], [], []
    chunk_index = 0
    for block in blocks:
        chunks = chunk_text(block["content"], config.CHUNK_SIZE, config.CHUNK_OVERLAP)
        for chunk in chunks:
            embedding = embed_text(chunk)
            all_chunks.append(chunk)
            all_embeddings.append(embedding)
            all_metadatas.append(
                {
                    "doc_id": doc_id,
                    "chat_id": chat_id,
                    "filename": filename,
                    "chunk_index": chunk_index,
                    "source_type": "text",
                    "page": block["page"],  # int for PDF, None for DOCX/TXT
                }
            )
            all_ids.append(f"{doc_id}_chunk_{chunk_index}")
            chunk_index += 1

    vector_store.add_chunks(all_chunks, all_embeddings, all_metadatas, all_ids)
    return {
        "doc_id": doc_id,
        "filename": filename,
        "chunk_count": chunk_index,
        "status": "ok",
    }


def parse_pdf(file_path: str) -> list[dict]:
    """Chunk per page so page numbers are accurate.
    Returns [{content: str, page: int}] — one entry per non-empty page."""
    doc = fitz.open(file_path)
    blocks = []
    for page_num, page in enumerate(doc):
        text = page.get_text().strip()
        if text:
            blocks.append({"content": text, "page": page_num + 1})  # 1-indexed
    doc.close()
    return blocks


def parse_docx(file_path: str) -> list[dict]:
    """Extract paragraphs and table cells as one block. page is None."""
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return [{"content": "\n".join(parts), "page": None}]


def parse_txt(file_path: str) -> list[dict]:
    """Read as UTF-8. Returns [{content: str, page: None}]."""
    text = Path(file_path).read_text(encoding="utf-8")
    return [{"content": text, "page": None}]
